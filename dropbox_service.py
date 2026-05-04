"""
Integración Dropbox para Mollo.
Soporta: listar, buscar, descargar, subir y analizar archivos.
Formatos soportados para análisis: PDF, DOCX, XLSX, TXT, CSV, MD.
"""
import io
import os
import tempfile
from pathlib import Path
from typing import Optional

import dropbox
from dropbox.exceptions import AuthError, ApiError
from dropbox.files import FileMetadata, FolderMetadata, SearchMatchV2

_client: Optional[dropbox.Dropbox] = None


def get_client() -> dropbox.Dropbox:
    global _client
    if _client:
        return _client
    # Leer directo del .env para evitar problemas de caching de os.environ
    from dotenv import dotenv_values
    from pathlib import Path
    env_path = Path(__file__).parent / ".env"
    env = dotenv_values(env_path)
    app_key       = env.get("DROPBOX_APP_KEY", "")
    app_secret    = env.get("DROPBOX_APP_SECRET", "")
    refresh_token = env.get("DROPBOX_REFRESH_TOKEN", "")
    if not refresh_token:
        raise RuntimeError(
            "DROPBOX_REFRESH_TOKEN no configurado. "
            "Ejecuta: python dropbox_setup.py"
        )
    _client = dropbox.Dropbox(
        app_key=app_key,
        app_secret=app_secret,
        oauth2_refresh_token=refresh_token,
    )
    return _client


# ── Listar archivos ───────────────────────────────────────────────────────────

def listar_archivos(carpeta: str = "") -> str:
    dbx = get_client()
    path = "" if carpeta in ("", "/", "raiz") else f"/{carpeta.lstrip('/')}"
    try:
        result = dbx.files_list_folder(path, limit=50)
        entries = result.entries
        while result.has_more:
            result = dbx.files_list_folder_continue(result.cursor)
            entries.extend(result.entries)
    except ApiError as e:
        return f"Error al listar '{path}': {e}"

    if not entries:
        return f"La carpeta '{path or 'raíz'}' está vacía."

    lines = [f"Contenido de {path or 'raíz'} ({len(entries)} elementos):"]
    carpetas = [e for e in entries if isinstance(e, FolderMetadata)]
    archivos = [e for e in entries if isinstance(e, FileMetadata)]

    for c in sorted(carpetas, key=lambda x: x.name):
        lines.append(f"  📁 {c.name}/")
    for a in sorted(archivos, key=lambda x: x.name):
        size = f"{a.size / 1024:.1f} KB" if a.size < 1024 * 1024 else f"{a.size / 1024 / 1024:.1f} MB"
        mod  = a.client_modified.strftime("%d/%m/%Y") if a.client_modified else ""
        lines.append(f"  📄 {a.name}  ({size}, {mod})")

    return "\n".join(lines)


# ── Buscar archivos ───────────────────────────────────────────────────────────

def buscar_archivos(query: str) -> str:
    dbx = get_client()
    try:
        result = dbx.files_search_v2(query, options=dropbox.files.SearchOptions(max_results=15))
    except ApiError as e:
        return f"Error en búsqueda: {e}"

    matches = [m for m in result.matches if isinstance(m, SearchMatchV2)]
    if not matches:
        return f"No se encontraron archivos con '{query}'."

    lines = [f"Resultados para '{query}':"]
    for m in matches:
        meta = m.metadata.get_metadata()
        if isinstance(meta, FileMetadata):
            lines.append(f"  📄 {meta.path_display}  ({meta.size / 1024:.1f} KB)")
        elif isinstance(meta, FolderMetadata):
            lines.append(f"  📁 {meta.path_display}/")
    return "\n".join(lines)


# ── Descargar y extraer texto ─────────────────────────────────────────────────

def _extraer_texto(data: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()

    if ext in (".txt", ".md", ".csv"):
        return data.decode("utf-8", errors="replace")[:15000]

    if ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [p.extract_text() or "" for p in pdf.pages[:20]]
        return "\n\n".join(pages)[:15000]

    if ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return text[:15000]

    if ext in (".xlsx", ".xls"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets[:3]:
            lines.append(f"[Hoja: {sheet.title}]")
            for row in sheet.iter_rows(max_row=200, values_only=True):
                row_text = "\t".join(str(c) if c is not None else "" for c in row)
                if row_text.strip():
                    lines.append(row_text)
        return "\n".join(lines)[:15000]

    return f"Formato '{ext}' no soportado para extracción de texto."


def descargar_texto(ruta: str) -> tuple[str, str]:
    """Devuelve (texto_extraido, nombre_archivo)."""
    dbx  = get_client()
    path = f"/{ruta.lstrip('/')}"
    try:
        meta, response = dbx.files_download(path)
        data = response.content
        texto = _extraer_texto(data, meta.name)
        return texto, meta.name
    except ApiError as e:
        return f"Error descargando '{path}': {e}", ""


# ── Subir archivo ─────────────────────────────────────────────────────────────

def subir_archivo(ruta_local: str, destino_dropbox: str) -> str:
    dbx = get_client()
    ruta_local = ruta_local.strip()
    if not os.path.exists(ruta_local):
        return f"Archivo local no encontrado: {ruta_local}"

    dest = f"/{destino_dropbox.lstrip('/')}"
    with open(ruta_local, "rb") as f:
        data = f.read()

    size = len(data)
    try:
        if size <= 150 * 1024 * 1024:  # ≤150 MB → upload simple
            meta = dbx.files_upload(
                data, dest,
                mode=dropbox.files.WriteMode.overwrite,
            )
        else:
            return "Archivo demasiado grande (>150 MB). Usa el cliente de Dropbox directamente."
        return f"Subido: {meta.path_display} ({size / 1024:.1f} KB)"
    except ApiError as e:
        return f"Error subiendo a Dropbox: {e}"


# ── Subir bytes directamente ──────────────────────────────────────────────────

def subir_bytes(data: bytes, filename: str, ruta_dropbox: str) -> str:
    """Sube bytes directamente a Dropbox sin necesitar archivo en disco."""
    dbx = get_client()
    dest = f"/{ruta_dropbox.lstrip('/')}"
    try:
        if len(data) > 150 * 1024 * 1024:
            return "Archivo demasiado grande (>150 MB)."
        meta = dbx.files_upload(data, dest, mode=dropbox.files.WriteMode.overwrite)
        return f"Subido: {meta.path_display} ({len(data) / 1024:.1f} KB)"
    except ApiError as e:
        return f"Error subiendo a Dropbox: {e}"


# ── Descargar archivo a disco ─────────────────────────────────────────────────

def descargar_archivo(ruta_dropbox: str, destino_local: str) -> str:
    dbx  = get_client()
    path = f"/{ruta_dropbox.lstrip('/')}"
    try:
        meta, response = dbx.files_download(path)
        dest = os.path.expanduser(destino_local)
        os.makedirs(os.path.dirname(dest) or ".", exist_ok=True)
        with open(dest, "wb") as f:
            f.write(response.content)
        return f"Descargado: {meta.name} → {dest} ({meta.size / 1024:.1f} KB)"
    except ApiError as e:
        return f"Error descargando '{path}': {e}"
